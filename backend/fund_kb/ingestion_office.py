"""OOXML adapters: preserve source-part/cell anchors, never run Office code."""
from __future__ import annotations

import io
import re
import zipfile
from datetime import date, datetime
from xml.etree import ElementTree as ET

from .ingestion import (
    ParsedDocument,
    ParseError,
    append_text,
    make_block,
    preview_document,
    render_blocks,
    safe_xml,
)

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _word_text(node: ET.Element, warnings: list[str]) -> str:
    if node.tag in {W + "del", W + "moveFrom"}:
        warnings.append("TRACKED_CHANGES_REQUIRE_REVIEW")
        return ""
    if node.tag in {W + "ins", W + "moveTo"}:
        warnings.append("TRACKED_CHANGES_REQUIRE_REVIEW")
    if node.tag == W + "r" and any(
        n.tag in {W + "vanish", W + "webHidden"} and n.get(W + "val", "1") not in {"0", "false", "off"}
        for n in node.iter()
    ):
        warnings.append("HIDDEN_TEXT_EXCLUDED")
        return ""
    if node.tag in {W + "t", "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"}:
        return node.text or ""
    if node.tag == W + "tab":
        return "\t"
    if node.tag in {W + "br", W + "cr"}:
        return "\n"
    if node.tag in {W + "footnoteReference", W + "endnoteReference"}:
        return f"[注{node.get(W + 'id', '')}]"
    return "".join(_word_text(c, warnings) for c in node)


def parse_docx(data: bytes, sha: str) -> ParsedDocument:
    blocks: list[dict] = []
    warnings = ["DOCX_REFLOW_NOT_NATIVE_PAGINATION"]
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        # Main text first; headers/footers/notes remain separately labelled evidence.
        parts = ["word/document.xml"] + sorted(n for n in archive.namelist()
                    if re.fullmatch(r"word/(header\d+|footer\d+|footnotes|endnotes)\.xml", n))
        for part in parts:
            root = safe_xml(archive.read(part))

            def walk(node: ET.Element, path: str, part=part):
                loc = {"kind": "docx", "part": part, "node_path": path,
                       "label": f"{part} · {path}"}
                if node.tag in {W + "del", W + "moveFrom"}:
                    warnings.append("TRACKED_CHANGES_REQUIRE_REVIEW")
                    return
                if node.tag in {W + "ins", W + "moveTo"}:
                    warnings.append("TRACKED_CHANGES_REQUIRE_REVIEW")
                if node.tag in {W + "footnote", W + "endnote"} and int(node.get(W + "id", "0")) <= 0:
                    return
                if node.tag == W + "p":
                    text = _word_text(node, warnings)
                    para_id = node.get("{http://schemas.microsoft.com/office/word/2010/wordml}paraId")
                    if para_id:
                        loc["paragraph_id"] = para_id
                    style = node.find("./" + W + "pPr/" + W + "pStyle")
                    style_id = style.get(W + "val", "") if style is not None else ""
                    heading = re.search(r"(?:Heading|标题)([1-6])", style_id, re.IGNORECASE)
                    append_text(blocks, text, loc, sha, "heading" if heading else "paragraph",
                                {"level": int(heading.group(1))} if heading else None)
                    if any(n.tag in {W + "drawing", W + "pict", W + "object"} for n in node.iter()):
                        warnings.append("DOCX_DRAWING_REQUIRES_ORIGINAL_REVIEW")
                elif node.tag == W + "tbl":
                    rows, cell_paths = [], []
                    for r, tr in enumerate(node.findall(W + "tr"), 1):
                        row = []
                        for c, tc in enumerate(tr.findall(W + "tc"), 1):
                            row.append("\n".join(_word_text(p, warnings) for p in tc.findall(W + "p")))
                            cell_paths.append({"row": r, "column": c, "node_path": f"{path}/tr[{r}]/tc[{c}]"})
                            if tc.find(".//" + W + "gridSpan") is not None or tc.find(".//" + W + "vMerge") is not None:
                                warnings.append("MERGED_TABLE_REFLOW_REQUIRES_REVIEW")
                            for n, nested in enumerate(tc.findall(W + "tbl"), 1):
                                walk(nested, f"{path}/tr[{r}]/tc[{c}]/tbl[{n}]")
                        rows.append(row)
                    if rows:
                        width = max(map(len, rows))
                        if width > 100 or len(rows) > 2000:
                            raise ParseError("TABLE_LIMIT", "Word表格超出限制。")
                        rows = [row + [""] * (width - len(row)) for row in rows]
                        # Do not invent that first row is a semantic header.
                        loc["cells"] = cell_paths
                        blocks.append(make_block("table", {"columns": [f"列{i + 1}" for i in range(width)],
                                                             "rows": rows}, loc, sha, len(blocks)))
                else:
                    counts: dict[str, int] = {}
                    for child in node:
                        name = child.tag.rsplit("}", 1)[-1]
                        counts[name] = counts.get(name, 0) + 1
                        walk(child, f"{path}/{name}[{counts[name]}]")
            walk(root, "/" + root.tag.rsplit("}", 1)[-1] + "[1]")
    return ParsedDocument(blocks, preview_document(render_blocks(blocks),
                          "Word结构重排预览；段落、表格与脚注定位以原始OOXML为准，不代表Word原生分页。"),
                          sorted(set(warnings)), {"preview_kind": "DOCX_REFLOW", "source_parts": parts})


def _scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def parse_xlsx(data: bytes, sha: str) -> ParsedDocument:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.cell import coordinate_to_tuple

    blocks: list[dict] = []
    warnings: list[str] = []
    # Two independent read-only views; openpyxl never recalculates formulas.
    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    values = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    sheet_details = []
    try:
        for ws in formulas.worksheets:
            sheet_details.append({"name": ws.title, "state": ws.sheet_state,
                                  "rows": ws.max_row, "columns": ws.max_column})
            if ws.sheet_state != "visible":
                warnings.append("HIDDEN_SHEET_EXCLUDED")
                continue
            # Formatting-only cells often inflate a vendor sheet to a million
            # rows. Validate actual populated coordinates, including coordinates
            # outside a dishonest small declared dimension, before iterating.
            max_row = max_col = 0
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                part = ws._worksheet_path
                if not re.fullmatch(r"xl/worksheets/[^/]+\.xml", part):
                    raise ParseError("INVALID_XLSX", "工作表路径无效")
                sheet_xml = safe_xml(archive.read(part))
                for cell in sheet_xml.iter(S + "c"):
                    if not any(child.tag in {S + "v", S + "f", S + "is"} for child in cell):
                        continue
                    try:
                        row_no, col_no = coordinate_to_tuple(cell.get("r", ""))
                    except (ValueError, TypeError, KeyError):
                        raise ParseError("INVALID_XLSX", "单元格坐标无效") from None
                    if row_no > 10000 or col_no > 100:
                        raise ParseError("SHEET_LIMIT", "实际内容超过10000行或100列；请拆分。")
                    max_row, max_col = max(max_row, row_no), max(max_col, col_no)
            if (max_row, max_col) != (ws.max_row or 0, ws.max_column or 0):
                warnings.append("FORMATTING_ONLY_DIMENSION_TRIMMED")
            sheet_details[-1].update(parsed_rows=max_row, parsed_columns=max_col)
            if not max_row or not max_col:
                continue
            cached_ws = values[ws.title]
            bounds = {"max_row": max_row, "max_col": max_col}
            for row_no, (row, cached_row) in enumerate(zip(ws.iter_rows(**bounds), cached_ws.iter_rows(**bounds), strict=True), 1):
                if row_no > 10000 or len(row) > 100:
                    raise ParseError("SHEET_LIMIT", "工作表展开超出限制。")
                if not any(c.value is not None for c in row):
                    continue
                details, row_text, formula_blocks = {}, [], []
                for i, (cell, cached) in enumerate(zip(row, cached_row, strict=True), 1):
                    coord = f"{get_column_letter(i)}{row_no}"
                    value = _scalar(cell.value)
                    info = {"value": value, "number_format": cell.number_format,
                            "data_type": cell.data_type}
                    if cell.data_type == "f":
                        expression = cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", str(cell.value))
                        cached_value = _scalar(cached.value)
                        info = {"formula": expression, "cached_value": cached_value,
                                "calculated_value": None, "value_status": "CACHED_UNVERIFIED" if cached_value is not None else "CACHE_MISSING",
                                "number_format": cell.number_format}
                        warnings.append("FORMULA_CACHE_UNVERIFIED" if cached_value is not None else "FORMULA_CACHE_MISSING")
                        value = f"{expression} [缓存{'未提供' if cached_value is None else '未核验: ' + str(cached_value)}]"
                        formula_blocks.append((coord, expression, info))
                    details[coord] = info
                    row_text.append("" if value is None else str(value))
                cell_range = f"A{row_no}:{get_column_letter(len(row))}{row_no}"
                loc = {"kind": "xlsx", "sheet": ws.title, "cell": cell_range, "cells": details,
                       "label": f"{ws.title}!{cell_range}"}
                blocks.append(make_block("table", {"columns": [get_column_letter(i) for i in range(1, len(row) + 1)],
                                                     "rows": [row_text]}, loc, sha, len(blocks)))
                for coord, expression, info in formula_blocks:
                    blocks.append(make_block("formula", {"expression_text": expression, "unit": "UNKNOWN",
                                                            "calculator_ref": None},
                                             {"kind": "xlsx", "sheet": ws.title, "cell": coord,
                                              "label": f"{ws.title}!{coord}", **info}, sha, len(blocks)))
                if len(blocks) > 10000:
                    raise ParseError("DOCUMENT_LIMIT", "Excel解析块数超出限制。")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            merges = {}
            for name in archive.namelist():
                if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name):
                    root = safe_xml(archive.read(name))
                    found = [n.get("ref") for n in root.iter(S + "mergeCell")]
                    if found:
                        merges[name] = found
                        warnings.append("MERGED_CELLS_ANCHORED_TO_TOP_LEFT")
                    if any(n.get("hidden") == "1" for n in root.iter() if n.tag in {S + "row", S + "col"}):
                        warnings.append("HIDDEN_ROWS_OR_COLUMNS_REQUIRE_REVIEW")
    finally:
        formulas.close()
        values.close()
    return ParsedDocument(blocks, preview_document(render_blocks(blocks),
                          "Excel按单元格只读预览；公式未执行，缓存值不代表重新计算或业务核验。"),
                          sorted(set(warnings)), {"preview_kind": "XLSX_CELL_GRID", "sheets": sheet_details,
                                                  "merged_ranges": merges, "formula_recalculated": False})
