"""Idempotent, explicit development corpus. Never runs against production."""
from __future__ import annotations

import hashlib
import io
import json
from datetime import date
from uuid import NAMESPACE_URL, uuid5

from sqlalchemy import select

from .db import utcnow


def uid(number: int) -> str:
    return f"00000000-0000-4000-8000-{number:012d}"


def sha(value: object) -> str:
    data = value if isinstance(value, bytes) else json.dumps(value, ensure_ascii=False, sort_keys=True).encode()
    return hashlib.sha256(data).hexdigest()


DISCLAIMER = "本资料为系统演示语料，不是任何基金公司的正式制度或现行法规。请上传并审核实际业务资料。"
DOCUMENTS = [
    (201, "基金估值业务操作规程", "docx", "估值与核算", True, [
        ("核对适用范围", "确认基金、份额类别、业务日期和估值处理状态，取得适用制度和数据。"),
        ("核对输入版本", "比较管理人与托管方的持仓、价格、费用参数及数据文件版本，记录不一致项目。"),
        ("复核费用计算", "核对费用基数、费率、计费期间和舍入口径，使用已确认的输入进行独立复核。"),
        ("提交复核材料", "形成差异原因、引用资料和未解决事项清单，由相应复核岗位按既有程序处理。"),
    ]),
    (202, "开放式基金份额登记指南", "pdf", "登记结算", True, [
        ("识别业务类型", "区分申购、赎回、转换、红利再投资和份额拆分，不能仅根据份额变化套用同一流程。"),
        ("检查日期与批次", "分别核对申请、确认、计价和支付日期；对持有期问题补充份额持有批次。"),
        ("核对份额变化", "根据来源记录核对期初份额、增加、减少和期末份额；重处理需核对原业务事件标识。"),
    ]),
    (203, "净值差异排查清单", "xlsx", "估值与核算", False, [
        ("数量差异", "核对持仓和交易批次"), ("价格差异", "核对估值日期、价格来源和权利限制"),
        ("费用差异", "核对基数、费率、期间和精度"), ("份额差异", "核对TA确认和业务事件"),
    ]),
    (204, "定期报告编制与复核", "docx", "报告管理", True, [
        ("选择报告版本", "记录报告期、产品类型、采用的规则与模板版本，历史报告按当时适用版本复核。"),
        ("核对数据来源", "每项关键指标记录来源、币种、单位、期间和计算口径，避免含义相近的指标混用。"),
        ("核对披露状态", "区分内部编制、专业复核、外部确认和正式披露，不以文档保存成功证明披露完成。"),
    ]),
    (205, "费用计提说明", "pdf", "费用管理", False, [
        ("确认承担主体", "分别识别基金财产、基金管理人和投资者承担的费用。"),
        ("确认适用版本", "结合业务日期、份额类别、有效合同和参数版本核对费用口径。"),
    ]),
]


def make_file(title: str, extension: str, sections: list[tuple[str, str]]) -> bytes:
    stream = io.BytesIO()
    if extension == "docx":
        from docx import Document
        document = Document()
        document.add_heading(title, 0)
        document.add_paragraph(DISCLAIMER)
        for heading, body in sections:
            document.add_heading(heading, 1)
            document.add_paragraph(body)
        document.save(stream)
    elif extension == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        book = Workbook()
        sheet = book.active
        sheet.title = "排查清单"
        sheet.append(["核对项目", "检查方法", "状态"])
        for heading, body in sections:
            sheet.append([heading, body, "待核对"])
        sheet.append(["资料说明", DISCLAIMER, "演示"])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="315FD4")
        sheet.column_dimensions["A"].width = 20
        sheet.column_dimensions["B"].width = 70
        sheet.column_dimensions["C"].width = 16
        book.save(stream)
    else:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        body_style = ParagraphStyle("body", fontName="STSong-Light", fontSize=11, leading=21,
                                    wordWrap="CJK", textColor="#26344D")
        title_style = ParagraphStyle("title", parent=body_style, fontSize=20, leading=30, spaceAfter=16)
        heading_style = ParagraphStyle("heading", parent=body_style, fontSize=14, leading=23, spaceBefore=16)
        story = [Paragraph(title, title_style), Paragraph(DISCLAIMER, body_style), Spacer(1, 16)]
        for heading, body in sections:
            story.extend([Paragraph(heading, heading_style), Paragraph(body, body_style)])
        SimpleDocTemplate(stream, pagesize=A4, leftMargin=48, rightMargin=48,
                          topMargin=48, bottomMargin=48).build(story)
    return stream.getvalue()


def seed_demo(session_factory, settings, vector_index=None) -> dict:
    if settings.app_env != "development" or settings.auth_mode != "demo":
        return {"seeded": False, "reason": "not_development_demo"}
    from .ingestion import block_text, parse_file, render_blocks, text_sha256
    from .models import (
        Blob,
        ContentBlock,
        EvidenceLink,
        RelationEdge,
        Release,
        Resource,
        ResourceVersion,
        ReviewDecision,
        Space,
        SpaceMember,
        User,
    )
    from .storage import Storage
    storage = Storage(settings)
    counts = {"resources": 0, "knowledge": 0}
    with session_factory() as session:
        for number, name, roles in [(1, "李明", ["reader", "editor", "admin"]),
                                    (2, "张慧", ["reader", "reviewer", "publisher"]),
                                    (3, "王磊", ["reader"])]:
            if not session.get(User, uid(number)):
                session.add(User(id=uid(number), external_subject="demo:admin" if number == 1 else f"demo:{number}", display_name=name,
                                 active=True))
        if not session.get(Space, uid(101)):
            session.add(Space(id=uid(101), name="基金运营部", revision=1))
        session.flush()
        for number, roles in [(1, ["reader", "editor", "admin"]),
                              (2, ["reader", "reviewer", "publisher"]), (3, ["reader"])]:
            for role in roles:
                if not session.get(SpaceMember, (uid(101), uid(number), role)):
                    session.add(SpaceMember(space_id=uid(101), user_id=uid(number), role=role))
        session.commit()

        for number, title, extension, category, published, sections in DOCUMENTS:
            if session.get(Resource, uid(number)):
                continue
            version_id, blob_id = uid(number + 1000), uid(number + 2000)
            raw = make_file(title, extension, sections)
            filename = f"{title}.{extension}"
            key = f"blobs/{blob_id}/{filename}"
            storage.write_bytes(key, raw)
            parsed = parse_file(storage.local_path(key), filename)
            blocks = parsed.blocks
            mime = {"pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}[extension]
            resource = Resource(id=uid(number), space_id=uid(101), kind="document", name=title,
                                category=category, tags=["演示资料", "运营知识"], owner_id=uid(1),
                                classification="INTERNAL", revision=1, access_epoch=1,
                                restricted=False, suspended=False, legal_hold=False)
            session.add(resource)
            session.add(Blob(id=blob_id, space_id=uid(101), object_key=key, sha256=sha(raw),
                             size_bytes=len(raw), mime_type=mime, scan_state="CLEAN"))
            session.flush()
            version = ResourceVersion(id=version_id, resource_id=resource.id, version_no=1,
                                      state="APPROVED" if published else "IN_REVIEW", author_id=uid(1),
                                      title=title, knowledge_type="source", origin="UPLOAD",
                                      source_blob_id=blob_id, source_verified=published,
                                      legal_status="NOT_APPLICABLE", valid_from=date(2026, 1, 1),
                                      applicability={}, required_facts=[], content_sha256=sha(blocks),
                                      change_kind="UPDATE", change_reason="创建合成演示资料", revision=1)
            session.add(version)
            session.flush()
            for ordinal, block in enumerate(blocks):
                session.add(ContentBlock(version_id=version_id, block_id=block["block_id"], ordinal=ordinal,
                                         block_type=block["block_type"], data=block["data"],
                                         search_text=block_text(block),
                                         locator=block.get("locator", {}), content_sha256=text_sha256(block_text(block))))
            storage.write_bytes(f"previews/{version_id}.html", render_blocks(blocks, "html").encode())
            storage.write_bytes(f"previews/{version_id}.md", render_blocks(blocks, "markdown").encode())
            from .services import check_frozen_hash
            session.flush()
            version.content_sha256 = check_frozen_hash(session, version)
            if published:
                _release_seed(session, resource, version, ReviewDecision, Release)
            session.commit()
            counts["resources"] += 1

        # Knowledge is deliberately authored from a synthetic source, never presented as regulation.
        source_id = uid(1201)
        source_blocks = session.scalars(select(ContentBlock).where(ContentBlock.version_id == source_id)
                                        .order_by(ContentBlock.ordinal)).all()
        for number, title, kind, sections in [
            (301, "费用差异排查与复核", "sop", DOCUMENTS[0][5]),
            (302, "如何处理净值差异", "faq", [("核对范围", "首先确认同一基金、同一份额类别、同一业务日期和计算批次。"),
                                             ("定位原因", "核对数量、价格、费用、份额和时间口径，记录差异及证据。")]),
            (303, "运营问题解决方案模板", "solution_template", [("明确问题", "描述目标、已知事实和需要补齐的资料。"),
                                                                  ("形成处理步骤", "每步写明岗位、输入、动作、输出、核对方法和依据。")]),
        ]:
            if session.get(Resource, uid(number)):
                continue
            resource = Resource(id=uid(number), space_id=uid(101), kind="template" if number == 303 else "knowledge",
                                name=title, category="估值与核算", tags=["演示资料", "SOP"], owner_id=uid(1),
                                classification="INTERNAL", restricted=False, revision=1, access_epoch=1,
                                suspended=False, legal_hold=False)
            session.add(resource)
            session.flush()
            version_id = uid(number + 1000)
            payload = []
            for index, (heading, body) in enumerate(sections):
                payload.append({"block_id": str(uuid5(NAMESPACE_URL, f"demo:{number}:{index}")),
                                "ordinal": index, "block_type": "step", "data": {
                                    "action": f"{heading}：{body}", "owner_role": "基金会计 / 复核岗",
                                    "output": "检查记录与待处理事项", "verification": "逐项核对来源、日期和版本，记录不一致项目"},
                                "locator": {"label": f"处理步骤 {index + 1}"}, "citations": []})
            version = ResourceVersion(id=version_id, resource_id=resource.id, version_no=1, state="APPROVED",
                                      author_id=uid(1), title=title, knowledge_type=kind, origin="HUMAN",
                                      source_verified=True, legal_status="NOT_APPLICABLE",
                                      valid_from=date(2026, 1, 1), applicability={}, required_facts=[],
                                      content_sha256=sha(payload), change_reason=DISCLAIMER, change_kind="UPDATE",
                                      revision=1)
            session.add(version)
            session.flush()
            for block in payload:
                session.add(ContentBlock(version_id=version_id, block_id=block["block_id"],
                                         ordinal=block["ordinal"], block_type="step", data=block["data"],
                                         search_text=block_text(block), locator=block["locator"],
                                         content_sha256=text_sha256(block_text(block))))
            session.flush()
            if source_blocks and number != 303:
                for index, block in enumerate(payload):
                    target = source_blocks[min(index + 1, len(source_blocks) - 1)]
                    session.add(EvidenceLink(id=str(uuid5(NAMESPACE_URL, f"link:{number}:{index}")),
                                             from_version_id=version_id, from_block_id=block["block_id"],
                                             to_version_id=source_id, to_block_id=target.block_id,
                                             purpose="INTERNAL_OPINION"))
                session.add(RelationEdge(id=str(uuid5(NAMESPACE_URL, f"relation:{number}")),
                                         source_version_id=version_id, target_resource_id=uid(201),
                                         relation_type="EXPLAINS", conditions={}))
            storage.write_bytes(f"previews/{version_id}.html", render_blocks(payload, "html").encode())
            storage.write_bytes(f"previews/{version_id}.md", render_blocks(payload, "markdown").encode())
            _release_seed(session, resource, version, ReviewDecision, Release)
            session.commit()
            counts["knowledge"] += 1
    return {"seeded": True, **counts, "corpus": "synthetic-development-only"}


def _release_seed(session, resource, version, ReviewDecision, Release):
    from .services import check_frozen_hash
    session.flush()
    version.content_sha256 = check_frozen_hash(session, version)
    session.flush()
    session.add(ReviewDecision(id=str(uuid5(NAMESPACE_URL, f"review:{version.id}")),
                               version_id=version.id, reviewer_id=uid(2), decision="APPROVE",
                               reviewed_sha256=version.content_sha256, comment="合成演示资料，非机构正式制度"))
    release = Release(id=str(uuid5(NAMESPACE_URL, f"release:{version.id}")), resource_id=resource.id,
                      version_id=version.id, state="ACTIVE", publisher_id=uid(2), activated_at=utcnow(),
                      manifest={"corpus": "synthetic-development-only", "version_id": version.id,
                                "content_sha256": version.content_sha256})
    session.add(release)
    session.flush()
    resource.active_release_id = release.id
