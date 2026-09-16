"""Regression cases discovered while importing the user-authorized corpus."""
import io

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from test_api import SPACE
from test_api import api as api  # noqa: PLC0414

from fund_kb import auth
from fund_kb.ingestion import ParseError, parse_file


def test_formatting_at_last_excel_row_does_not_fake_large_content(tmp_path):
    wb=Workbook();ws=wb.active;ws['A1']='指标';ws['B1']='释义';ws['A2']='估值净价'
    ws.cell(1048576,16).fill=PatternFill('solid',fgColor='FFFFFF')
    out=io.BytesIO();wb.save(out);path=tmp_path/'formatting.xlsx';path.write_bytes(out.getvalue())
    result=parse_file(path,path.name)
    assert len(result.blocks)==2
    assert 'FORMATTING_ONLY_DIMENSION_TRIMMED' in result.warnings
    assert result.metadata['sheets'][0]['parsed_rows']==2


def test_real_far_excel_content_still_exceeds_limit(tmp_path):
    wb=Workbook();wb.active.cell(10001,1,'实际数据')
    out=io.BytesIO();wb.save(out);path=tmp_path/'large.xlsx';path.write_bytes(out.getvalue())
    with pytest.raises(ParseError) as exc:
        parse_file(path,path.name)
    assert exc.value.code=='SHEET_LIMIT'


def test_body_is_buffered_before_database_authentication(api, monkeypatch):
    original=auth.authenticate
    seen=[]
    def check(request,db):
        assert isinstance(request.state.prefetched_request_body,bytes)
        seen.append(True)
        return original(request,db)
    monkeypatch.setattr(auth,'authenticate',check)
    response=api.call('POST','/resources',{'space_id':SPACE,'kind':'knowledge','name':'合成导入测试'})
    assert response.status_code==201 and seen


def test_buffering_never_replaces_auth_or_csrf_checks(api):
    response=api.client.post('/api/v1/resources',json={'space_id':SPACE,'kind':'knowledge','name':'不能越权'},
        headers={'Origin':'http://testserver'})
    assert response.status_code in {401,403}


def test_source_properties_update_does_not_replace_content_or_approve_source(api):
    resource=api.resource('document','来源属性测试')
    version=api.draft(resource)
    edited=api.edit(version,text='保留原稿内容')
    vid=edited.json()['id']
    result=api.call('PATCH',f'/document-versions/{vid}/source-properties',
        {'source_url':'https://example.org/source'},edited.headers['etag'])
    assert result.status_code==200,result.text
    after=api.call('GET','/versions/'+vid).json()
    assert after['blocks']==edited.json()['blocks']
    assert not after['source_verified'] and after['state']=='DRAFT'
    assert after['source_url']=='https://example.org/source'
    stale=api.call('PATCH',f'/document-versions/{vid}/source-properties',
        {'source_url':'https://example.org/other'},edited.headers['etag'])
    assert stale.status_code==412
