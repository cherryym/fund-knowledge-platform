from __future__ import annotations

import io
import zipfile
from html.parser import HTMLParser
from uuid import UUID

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

from fund_kb import ingestion_pdf
from fund_kb.ingestion import (
    ParseError,
    block_text,
    parse_file,
    render_blocks,
    safe_markdown,
    sanitize_html,
    text_sha256,
)


@pytest.fixture(autouse=True)
def no_ocr_model(monkeypatch):
    # All tests are offline; installed OCR engines must not accidentally run.
    monkeypatch.setattr(ingestion_pdf, "ocr_image", lambda image: ([], ["OCR_ENGINE_UNAVAILABLE"]))


def store(tmp_path, name, value):
    path = tmp_path / name
    path.write_bytes(value.encode("utf-8") if isinstance(value, str) else value)
    return path


def pdf_bytes(text="Fund operations control", encrypted=False, active=False, blank=False):
    writer = PdfWriter()
    page = writer.add_blank_page(300, 300)
    if not blank:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 30 230 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("synthetic-test-password")
    if active:
        writer.add_js("app.alert('test');")
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def office_bytes():
    doc = Document()
    doc.add_heading("基金资料核对", level=1)
    doc.add_paragraph("核对基金合同版本与业务日期。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "岗位"
    table.cell(0, 1).text = "核对"
    table.cell(1, 0).text = "运营复核岗"
    table.cell(1, 1).text = "核对份额类别"
    doc.sections[0].header.paragraphs[0].text = "合成资料页眉"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def rewrite_zip(data, extra=None, edit=None):
    target = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as out:
        for name in source.namelist():
            content = source.read(name)
            if edit:
                content = edit(name, content)
            out.writestr(name, content)
        for name, content in (extra or {}).items():
            out.writestr(name, content)
    return target.getvalue()


def assert_blocks(result):
    assert [b["ordinal"] for b in result.blocks] == list(range(len(result.blocks)))
    for block in result.blocks:
        UUID(block["block_id"])
        assert set(block) == {"block_id", "ordinal", "block_type", "data", "locator", "citations"}
        assert block["locator"]["label"]
    assert result.metadata["source_verified"] is False
    assert result.metadata["security_scan"]["antivirus"] == "NOT_PERFORMED"


def test_utf8_text_has_exact_lines_and_deterministic_ids(tmp_path):
    path = store(tmp_path, "资料.txt", "基金合同\n\n核对份额类别。\n")
    first, second = parse_file(path, path.name), parse_file(path, path.name)
    assert_blocks(first)
    assert first.blocks == second.blocks
    assert [b["locator"]["line_start"] for b in first.blocks] == [1, 3]
    assert first.blocks[1]["locator"]["source_char_start"] == len("基金合同\n\n")
    assert block_text(first.blocks[1]) == "核对份额类别。"


@pytest.mark.parametrize("encoding", ["utf-16", "utf-32", "gb18030"])
def test_text_encoding_is_explicit(tmp_path, encoding):
    path = store(tmp_path, "encoded.txt", "基金份额类别".encode(encoding))
    result = parse_file(path, path.name)
    assert "基金份额" in block_text(result.blocks[0])
    assert any("ENCODING" in w or "UTF8" in w for w in result.warnings)


def test_markdown_headings_lists_tables_and_source_maps(tmp_path):
    text = "# 合同核对\n\n- 核对日期\n- 核对类别\n\n| 岗位 | 输出 |\n| --- | --- |\n| 运营 | 清单 |\n\n```python\nprint('只是文本')\n```\n"
    result = parse_file(store(tmp_path, "sop.md", text), "sop.md")
    assert_blocks(result)
    assert [b["block_type"] for b in result.blocks] == ["heading", "list", "table", "paragraph"]
    assert result.blocks[1]["data"]["items"] == ["核对日期", "核对类别"]
    assert result.blocks[2]["data"]["rows"] == [["运营", "清单"]]
    assert result.blocks[0]["locator"]["line_start"] == 1
    assert "只是文本" in result.preview_html


@pytest.mark.parametrize("attack", [
    "<script>alert(1)</script><p>正常资料</p>",
    '<svg><script>alert(1)</script></svg><p onclick="alert(1)">正常资料</p>',
    '<a href="java&#x73;cript:alert(1)">打开</a><iframe src="https://example.invalid"></iframe>',
    '<img src="https://example.invalid/track" onerror="alert(1)"><style>body{display:none}</style>',
    '<object data="x"></object><form action="https://example.invalid"><input name="x"></form>',
])
def test_html_sanitizer_disallows_active_content(attack):
    cleaned = sanitize_html(attack)
    assert "<script" not in cleaned
    assert "<svg" not in cleaned
    assert "javascript:" not in cleaned
    assert "onerror=" not in cleaned
    assert "onclick=" not in cleaned
    assert "<img" not in cleaned
    assert "<iframe" not in cleaned
    assert "<form" not in cleaned


def test_markdown_render_never_fetches_remote_images_or_executes_raw_html():
    cleaned = safe_markdown('![x](https://example.invalid/track)\n<script>alert(1)</script>\n[x](javascript:alert(1))')
    assert "<img" not in cleaned
    assert "<script" not in cleaned
    assert 'href="javascript:' not in cleaned


def test_html_parses_safe_tables_and_nodes_not_script_or_hidden_text(tmp_path):
    source = '<html><body><h1>基金合同</h1>\n<p>核对日期</p><div hidden>隐藏草稿</div>' \
             '<script>隐藏指令</script><table><tr><th>项目</th><th>值</th></tr>' \
             '<tr><td>类别</td><td>A类</td></tr></table></body></html>'
    result = parse_file(store(tmp_path, "doc.html", source), "doc.html")
    assert_blocks(result)
    texts = "\n".join(block_text(b) for b in result.blocks)
    assert "隐藏指令" not in texts and "隐藏草稿" not in texts
    assert any(b["block_type"] == "table" for b in result.blocks)
    assert result.blocks[0]["locator"]["node_path"].endswith("/h1[1]")
    assert "HTML_ACTIVE_OR_HIDDEN_CONTENT_REMOVED" in result.warnings


def test_rendering_blocks_is_safe_in_both_formats():
    blocks = [{"block_id": "test", "block_type": "paragraph", "data": {"text": '<script>alert(1)</script> ![x](https://x.invalid)'}, "locator": {}}]
    assert "<script" not in render_blocks(blocks)
    md = render_blocks(blocks, "markdown")
    assert "<script" not in md
    assert "<img" not in safe_markdown(md)
    with pytest.raises(ValueError):
        render_blocks(blocks, "pdf")


def test_docx_keeps_body_table_header_order_and_structural_anchors(tmp_path):
    result = parse_file(store(tmp_path, "source.docx", office_bytes()), "source.docx")
    assert_blocks(result)
    assert result.blocks[0]["block_type"] == "heading"
    table = next(b for b in result.blocks if b["block_type"] == "table")
    assert table["data"]["rows"][1] == ["运营复核岗", "核对份额类别"]
    assert "tc[2]" in table["locator"]["cells"][1]["node_path"]
    assert any("header" in b["locator"]["part"] for b in result.blocks)
    assert all("source_page" not in b["locator"] for b in result.blocks)
    assert result.metadata["preview_kind"] == "DOCX_REFLOW"


def test_docx_footnotes_revisions_hidden_text_are_explicit(tmp_path):
    def change(name, content):
        if name == "word/document.xml":
            content = content.replace(b"</w:body>", b'<w:del><w:p><w:r><w:t>DELETED</w:t></w:r></w:p></w:del>'
                b'<w:p><w:r><w:rPr><w:vanish/></w:rPr><w:t>SECRET_HIDDEN</w:t></w:r></w:p>'
                b'<w:ins><w:p><w:r><w:t>INSERTED</w:t></w:r></w:p></w:ins></w:body>')
        return content
    footnotes = b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>FOOTNOTE</w:t></w:r></w:p></w:footnote></w:footnotes>'
    content = rewrite_zip(office_bytes(), {"word/footnotes.xml": footnotes}, change)
    result = parse_file(store(tmp_path, "review.docx", content), "review.docx")
    text = "\n".join(block_text(b) for b in result.blocks)
    assert "DELETED" not in text and "SECRET_HIDDEN" not in text
    assert "INSERTED" in text and "FOOTNOTE" in text
    assert "TRACKED_CHANGES_REQUIRE_REVIEW" in result.warnings
    assert "HIDDEN_TEXT_EXCLUDED" in result.warnings


def test_xlsx_values_formula_cache_and_cell_anchors(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "费用核对"
    ws.append(["项目", "金额"])
    ws.append(["管理费", 120])
    ws["B3"] = "=B2*2"
    wb.create_sheet("草稿").sheet_state = "hidden"
    buffer = io.BytesIO()
    wb.save(buffer)
    result = parse_file(store(tmp_path, "calc.xlsx", buffer.getvalue()), "calc.xlsx")
    assert_blocks(result)
    formula = next(b for b in result.blocks if b["block_type"] == "formula")
    assert formula["locator"]["sheet"] == "费用核对" and formula["locator"]["cell"] == "B3"
    assert formula["locator"]["cached_value"] is None
    assert formula["locator"]["calculated_value"] is None
    assert formula["data"]["expression_text"] == "=B2*2"
    assert "FORMULA_CACHE_MISSING" in result.warnings
    assert "HIDDEN_SHEET_EXCLUDED" in result.warnings
    assert result.metadata["formula_recalculated"] is False


def test_xlsx_cached_value_is_not_recalculated(tmp_path):
    wb = Workbook()
    wb.active["A1"] = "=1+1"
    buffer = io.BytesIO()
    wb.save(buffer)
    def edit(name, data):
        return data.replace(b"<v></v>", b"<v>999</v>") if name == "xl/worksheets/sheet1.xml" else data
    result = parse_file(store(tmp_path, "cache.xlsx", rewrite_zip(buffer.getvalue(), edit=edit)), "cache.xlsx")
    formula = next(b for b in result.blocks if b["block_type"] == "formula")
    assert formula["locator"]["cached_value"] == 999
    assert formula["locator"]["calculated_value"] is None
    assert "FORMULA_CACHE_UNVERIFIED" in result.warnings


def test_pdf_text_page_locator_and_security_metadata(tmp_path):
    result = parse_file(store(tmp_path, "source.pdf", pdf_bytes()), "source.pdf")
    assert_blocks(result)
    assert "Fund operations control" in " ".join(block_text(b) for b in result.blocks)
    assert result.blocks[0]["locator"]["source_page"] == 1
    assert result.metadata["page_count"] == 1
    assert result.metadata["parse_complete"] is True


def test_pdf_exact_bbox_and_raster_preview_when_dependencies_present(tmp_path):
    pytest.importorskip("pdfplumber")
    pytest.importorskip("pypdfium2")
    result = parse_file(store(tmp_path, "layout.pdf", pdf_bytes()), "layout.pdf")
    locator = result.blocks[0]["locator"]
    assert locator["coordinate_space"] == "top_left_points"
    assert 25 <= locator["bbox"][0] <= 35
    assert 40 < locator["bbox"][1] < 80
    assert "data:image/png;base64," in result.preview_html
    assert result.metadata["preview_kind"] == "PDF_RASTER"


@pytest.mark.parametrize("kwargs,code", [({"encrypted": True}, "ENCRYPTED_FILE"), ({"active": True}, "ACTIVE_CONTENT")])
def test_pdf_encryption_and_javascript_rejected(tmp_path, kwargs, code):
    with pytest.raises(ParseError) as err:
        parse_file(store(tmp_path, "bad.pdf", pdf_bytes(**kwargs)), "bad.pdf")
    assert err.value.code == code


def test_image_preview_remains_available_when_ocr_missing(tmp_path):
    image = Image.new("RGB", (100, 80), "white")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    result = parse_file(store(tmp_path, "scan.png", buffer.getvalue()), "scan.png")
    assert result.blocks == []
    assert "OCR_ENGINE_UNAVAILABLE" in result.warnings
    assert "NO_EXTRACTABLE_TEXT" in result.warnings
    assert "data:image/png;base64," in result.preview_html
    assert result.metadata["ocr_used"] is False


def test_ocr_adapter_word_boxes_map_to_source_image(monkeypatch, tmp_path):
    monkeypatch.setattr(ingestion_pdf, "ocr_image", lambda image: ([{"text": "基金合同", "bbox": [10, 20, 80, 40],
        "confidence": 88.0, "line": ["1", "1", "1"]}], ["OCR_TEXT_REQUIRES_REVIEW"]))
    buffer = io.BytesIO()
    Image.new("RGB", (100, 80), "white").save(buffer, format="PNG")
    result = parse_file(store(tmp_path, "ocr.png", buffer.getvalue()), "ocr.png")
    assert result.blocks[0]["locator"]["bbox"] == [10, 20, 80, 40]
    assert result.blocks[0]["locator"]["accuracy"] == "OCR_UNVERIFIED"
    assert result.metadata["source_verified"] is False


@pytest.mark.parametrize("extra,code", [
    ({"../escape.xml": "x"}, "UNSAFE_ZIP_PATH"),
    ({"word/vbaProject.bin": "x"}, "ACTIVE_CONTENT"),
    ({"word/activeX/activeX1.xml": "<x/>"}, "ACTIVE_CONTENT"),
    ({"word/embeddings/data.bin": "x"}, "ACTIVE_CONTENT"),
    ({"word/extra.xml": "<!DOCTYPE x [<!ENTITY e SYSTEM 'file:///sensitive'>]><x>&e;</x>"}, "UNSAFE_XML"),
    ({"word/large.xml": "<x>" + "A" * 100000 + "</x>"}, "ZIP_LIMIT"),
])
def test_ooxml_safety_checks(tmp_path, extra, code):
    with pytest.raises(ParseError) as err:
        parse_file(store(tmp_path, "bad.docx", rewrite_zip(office_bytes(), extra)), "bad.docx")
    assert err.value.code == code


def test_external_template_is_rejected(tmp_path):
    rels = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rX" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate" Target="https://example.invalid/t.dotx" TargetMode="External"/></Relationships>'
    with pytest.raises(ParseError) as err:
        parse_file(store(tmp_path, "external.docx", rewrite_zip(office_bytes(), {"word/_rels/settings.xml.rels": rels})), "external.docx")
    assert err.value.code == "EXTERNAL_REFERENCE"


@pytest.mark.parametrize("name,data,code", [
    ("x.txt", b"%PDF-1.7", "MIME_MISMATCH"),
    ("x.pdf", b"plain text", "MIME_MISMATCH"),
    ("x.doc", b"old-format", "UNSUPPORTED_FORMAT"),
    ("x.xlsm", b"macro", "ACTIVE_CONTENT"),
    ("x.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", "ENCRYPTED_OR_LEGACY_OFFICE"),
    ("x.txt", b"EICAR-STANDARD-ANTIVIRUS-TEST-FILE", "FILE_REJECTED"),
    ("x.txt", b"", "EMPTY_FILE"),
])
def test_invalid_formats_are_not_success(tmp_path, name, data, code):
    with pytest.raises(ParseError) as err:
        parse_file(store(tmp_path, name, data), name)
    assert err.value.code == code


class RenderedTags(HTMLParser):
    """Inspect actual elements/attributes, not inert escaped attack text."""

    def __init__(self, rendered):
        super().__init__(convert_charrefs=True)
        self.tags, self.attributes, self.text = [], [], []
        self.feed(rendered)
        self.close()

    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)
        self.attributes.extend(attrs)

    def handle_data(self, data):
        self.text.append(data)


def formatted_block(kind, text, **extra):
    return {"block_id": "00000000-0000-4000-8000-000000000009", "ordinal": 0,
        "block_type": kind, "data": {"text": text, **({"level": 3} if kind == "heading" else {}), **extra},
        "locator": {}, "citations": []}


@pytest.mark.parametrize("kind", ["heading", "paragraph", "warning"])
@pytest.mark.parametrize("explicit_plain", [False, True])
def test_legacy_and_explicit_plain_render_literal_in_html_and_markdown(kind, explicit_plain):
    text = '*字面量* **星号** _下划线_ [来源](https://example.invalid) <b>原文</b> &'
    block = formatted_block(kind, text, **({"text_format": "plain"} if explicit_plain else {}))
    html_view = RenderedTags(render_blocks([block], "html"))
    assert "".join(html_view.text) == text
    assert not set(html_view.tags) & {"strong", "em", "a", "b", "img"}
    exported = render_blocks([block], "markdown")
    assert r"\*字面量\*" in exported and r"\*\*星号\*\*" in exported
    markdown_view = RenderedTags(safe_markdown(exported))
    assert "".join(markdown_view.text).strip() == text
    assert not set(markdown_view.tags) & {"strong", "em", "a", "b", "img"}
    assert block_text(block) == text
    assert text_sha256(block_text(block)) == text_sha256(text)


@pytest.mark.parametrize("kind", ["heading", "paragraph", "warning"])
def test_explicit_markdown_renders_format_and_exports_canonical_text_once(kind):
    canonical = '**关键** 和 *说明* 与 `字段` [来源](https://example.invalid/rules)'
    block = formatted_block(kind, canonical, text_format="markdown")
    rendered = render_blocks([block], "html")
    assert "<strong>关键</strong>" in rendered and "<em>说明</em>" in rendered
    assert "<code>字段</code>" in rendered
    assert 'href="https://example.invalid/rules" rel="noreferrer noopener"' in rendered
    tags = RenderedTags(rendered).tags
    if kind == "heading":
        assert tags[0] == "h3" and "p" not in tags
    else:
        assert tags[:2] == ["section", "p"]
        assert "<p><p>" not in rendered
    prefix = "### " if kind == "heading" else ""
    assert render_blocks([block], "markdown") == prefix + canonical
    assert block_text(block) == canonical
    plain = formatted_block(kind, canonical)
    assert text_sha256(block_text(block)) == text_sha256(block_text(plain))


def test_marked_paragraph_preserves_paragraph_structure_and_heading_stays_inline():
    paragraph = render_blocks([formatted_block("paragraph", "**第一段**\n\n第二段  \n换行", text_format="markdown")])
    assert RenderedTags(paragraph).tags == ["section", "p", "strong", "p", "br"]
    heading = render_blocks([formatted_block("heading", "**标题**\n\n# 不得改变标题级别", text_format="markdown")])
    assert RenderedTags(heading).tags == ["h3", "strong"]


@pytest.mark.parametrize("attack", [
    '<script>alert(1)</script><b>原样HTML</b>',
    '<svg onload="alert(1)"><foreignObject>内容</foreignObject></svg>',
    '<iframe src="https://example.invalid/remote"></iframe>',
    '<img src="https://example.invalid/track" onerror="alert(1)">',
    '![远程图片](https://example.invalid/track)',
    '![嵌入图片](data:image/png;base64,eA==)',
    '[脚本](javascript:alert(1))',
    '[编码脚本](java&#x73;cript:alert(1))',
    '[数据链接](data:text/html;base64,PHNjcmlwdD4=)',
    '[文件链接](file:///private/file)',
    '[其他脚本](vbscript:msgbox(1))',
])
@pytest.mark.parametrize("kind", ["heading", "paragraph", "warning"])
def test_marked_block_html_is_sanitized_without_fetching_images(kind, attack):
    view = RenderedTags(render_blocks([formatted_block(kind, attack, text_format="markdown")]))
    assert not set(view.tags) & {"script", "svg", "math", "iframe", "img", "object", "embed", "style", "b"}
    for name, value in view.attributes:
        assert not name.startswith("on") and name not in {"src", "style"}
        if name == "href":
            assert value.startswith(("https://", "http://", "mailto:", "#"))


@pytest.mark.parametrize("marker", ["html", "MARKDOWN", "richtext", "", None, False, {}, []])
def test_render_rejects_unknown_or_non_string_markers(marker):
    with pytest.raises(ValueError, match="text_format"):
        render_blocks([formatted_block("paragraph", "内容", text_format=marker)])


def test_imported_text_is_not_implicitly_upgraded_to_markdown(tmp_path):
    raw = "**原始资料** [仅文字](https://example.invalid)\n*不能重解释*\n"
    path = store(tmp_path, "literal.txt", raw)
    parsed = parse_file(path, path.name)
    assert all("text_format" not in b["data"] for b in parsed.blocks)
    view = RenderedTags(parsed.preview_html)
    assert not set(view.tags) & {"strong", "em", "a"}
    assert "**原始资料**" in "".join(view.text)
    assert path.read_text() == raw
